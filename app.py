import zipfile
import streamlit as st
import cv2
import numpy as np
import os
import face_recognition
from openpyxl import Workbook, load_workbook
from PIL import Image
import tempfile
import onnxruntime as ort
import io
from datetime import datetime

# Paths
model_path = 'MODEL12/weights/best.onnx'
excel_path = 'face_detection_log.xlsx'
face_encodings_dict = {}

# Load model
@st.cache_resource
def load_model():
    return ort.InferenceSession(model_path)

session = load_model()

def create_excel_sheet():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Face Detection Log'
    sheet.append(['Face Name', 'Image Path'])
    workbook.save(excel_path)

def update_excel_sheet(face_name, image_path):
    if not os.path.exists(excel_path):
        create_excel_sheet()
    workbook = load_workbook(excel_path)
    sheet = workbook.active
    sheet.append([face_name, image_path])
    workbook.save(excel_path)

def preprocess(image):
    image = cv2.resize(image, (640, 640))
    image = image.transpose(2, 0, 1)
    image = np.expand_dims(image, axis=0).astype(np.float32) / 255.0
    return image

def non_max_suppression(boxes, scores, iou_threshold=0.4):
    indices = cv2.dnn.NMSBoxes(boxes, scores, 0.25, iou_threshold)
    return indices.flatten() if len(indices) > 0 else []

def postprocess(outputs, image_shape, input_shape):
    boxes, scores, class_ids = [], [], []
    for output in outputs:
        for det in output:
            if det[4] > 0.25:
                x_center, y_center, width, height = det[:4]
                left = int((x_center - width / 2) * (image_shape[1] / input_shape[3]))
                top = int((y_center - height / 2) * (image_shape[0] / input_shape[2]))
                right = int((x_center + width / 2) * (image_shape[1] / input_shape[3]))
                bottom = int((y_center + height / 2) * (image_shape[0] / input_shape[2]))
                boxes.append([left, top, right, bottom])
                scores.append(det[4])
                class_ids.append(int(det[5]))
    indices = non_max_suppression(boxes, scores)
    return [boxes[i] for i in indices], [scores[i] for i in indices], [class_ids[i] for i in indices]

def detect_faces(image):
    input_image = preprocess(image)
    outputs = session.run(None, {session.get_inputs()[0].name: input_image})
    boxes, scores, class_ids = postprocess(outputs[0], image.shape, input_image.shape)
    face_locations = [(box[1], box[2], box[3], box[0]) for box, cid in zip(boxes, class_ids) if cid == 0]
    rgb_image = cv2.cvtColor(image, cv2.COLOR_BGR2RGB)
    face_encodings = face_recognition.face_encodings(rgb_image, face_locations)
    return face_locations, face_encodings, image

def process_image(image, image_name, output_folder):
    global face_encodings_dict
    face_locations, face_encodings, original = detect_faces(image)
    for idx, (encoding, (top, right, bottom, left)) in enumerate(zip(face_encodings, face_locations)):
        match_found = False
        for name, encodings in face_encodings_dict.items():
            if True in face_recognition.compare_faces(encodings, encoding):
                match_found = True
                folder_path = os.path.join(output_folder, name)
                os.makedirs(folder_path, exist_ok=True)
                cv2.imwrite(os.path.join(folder_path, image_name), original)
                update_excel_sheet(name, os.path.join(folder_path, image_name))
                break
        if not match_found:
            face_image = original[top:bottom, left:right]
            if face_image is None or face_image.size == 0:
                st.warning(f"Skipped an empty or invalid face crop in image: {image_name}")
                continue
            face_pil = Image.fromarray(cv2.cvtColor(face_image, cv2.COLOR_BGR2RGB))
            st.image(face_pil, caption="New Face Detected")
            unique_key = f"{image_name}_{top}_{right}_{bottom}_{left}_{idx}"
            name = st.text_input(f"Enter name for face in {image_name} (face #{idx+1}):", key=unique_key)
        if name:
                face_encodings_dict[name] = [encoding]
                folder_path = os.path.join(output_folder, name)
                os.makedirs(folder_path, exist_ok=True)
                cv2.imwrite(os.path.join(folder_path, image_name), original)
                update_excel_sheet(name, os.path.join(folder_path, image_name))

st.title("📸 Face Detection & Organization")

option = st.radio("Choose Input Method", ["Upload ZIP Folder", "Capture via Webcam"])

if option == "Upload ZIP Folder":
    uploaded_zip = st.file_uploader("Upload a ZIP file with images", type="zip")
    if uploaded_zip:
        with tempfile.TemporaryDirectory() as tmpdir:
            zip_path = os.path.join(tmpdir, "images.zip")
            with open(zip_path, "wb") as f:
                f.write(uploaded_zip.read())
            with zipfile.ZipFile(zip_path, 'r') as zip_ref:
                zip_ref.extractall(tmpdir)
            output_dir = os.path.join(tmpdir, "sorted_faces")
            os.makedirs(output_dir, exist_ok=True)
            for root, dirs, files in os.walk(tmpdir):
                for file in files:
                    if file.lower().endswith(('.jpg', '.jpeg', '.png')):
                        img_path = os.path.join(root, file)
                        img = cv2.imread(img_path)
                        if img is not None:
                            process_image(img, file, output_dir)
            st.success("All images processed!")
            if os.path.exists(excel_path):
                with open(excel_path, "rb") as f:
                    st.download_button("Download Log File", f, file_name="face_detection_log.xlsx")

if option == "Capture via Webcam":
    frame = st.camera_input("Take a picture")
    if frame is not None:
        image = Image.open(frame)
        img_array = np.array(image.convert('RGB'))[:, :, ::-1]  # RGB to BGR
        temp_name = f"webcam_{datetime.now().strftime('%Y%m%d_%H%M%S')}.jpg"
        process_image(img_array, temp_name, "captured_faces")
        st.success("Webcam image processed!")
        if os.path.exists(excel_path):
            with open(excel_path, "rb") as f:
                st.download_button("Download Log File", f, file_name="face_detection_log.xlsx")
