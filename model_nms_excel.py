import cv2
import face_recognition
import numpy as np
import os
from shutil import copyfile
from openpyxl import Workbook, load_workbook
import onnxruntime as ort

# Path to the custom ONNX model
model_path = 'MODEL13/weights/best.onnx'
session = ort.InferenceSession(model_path)

# Excel file path
excel_path = 'face_detection_log.xlsx'

def create_excel_sheet():
    # Create a new Excel workbook and sheet if it doesn't exist
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Face Detection Log'
    sheet.append(['Face Name', 'Image Path'])
    workbook.save(excel_path)

def update_excel_sheet(face_name, image_path):
    # Load the existing workbook
    workbook = load_workbook(excel_path)
    sheet = workbook.active
    # Append the new face name and image path
    sheet.append([face_name, image_path])
    workbook.save(excel_path)

def preprocess(image):
    image = cv2.resize(image, (640, 640))
    image = image.transpose(2, 0, 1)  # HWC to CHW
    image = np.expand_dims(image, axis=0)
    image = image.astype(np.float32)
    image /= 255.0  # Normalize to [0, 1]
    return image

def non_max_suppression(boxes, scores, iou_threshold):
    indices = cv2.dnn.NMSBoxes(boxes, scores, score_threshold=0.25, nms_threshold=iou_threshold)
    return indices.flatten() if len(indices) > 0 else []

def postprocess(outputs, image_shape, input_shape, iou_threshold=0.4):
    boxes, scores, class_ids = [], [], []
    for output in outputs:
        for det in output:
            if det[4] > 0.25:  # Confidence threshold
                x_center, y_center, width, height = det[0], det[1], det[2], det[3]
                left = int((x_center - width / 2) * (image_shape[1] / input_shape[3]))
                top = int((y_center - height / 2) * (image_shape[0] / input_shape[2]))
                right = int((x_center + width / 2) * (image_shape[1] / input_shape[3]))
                bottom = int((y_center + height / 2) * (image_shape[0] / input_shape[2]))
                boxes.append([left, top, right, bottom])
                scores.append(det[4])
                class_ids.append(int(det[5]))

    indices = non_max_suppression(boxes, scores, iou_threshold)
    filtered_boxes = [boxes[i] for i in indices]
    filtered_scores = [scores[i] for i in indices]
    filtered_class_ids = [class_ids[i] for i in indices]

    return filtered_boxes, filtered_scores, filtered_class_ids

def detect_and_recognize_faces(image_path):
    if session is None:
        print("Model not loaded correctly. Exiting...")
        return [], [], None
    
    # Load the image
    image = cv2.imread(image_path)
    input_image = preprocess(image)
    
    # Run inference
    outputs = session.run(None, {session.get_inputs()[0].name: input_image})
    
    # Postprocess the outputs
    boxes, scores, class_ids = postprocess(outputs[0], image.shape, input_image.shape)
    
    # Filter only person class (assuming 'person' class index is 0)
    face_locations = []
    for box, score, class_id in zip(boxes, scores, class_ids):
        if class_id == 0:
            face_locations.append((box[1], box[2], box[3], box[0]))
    
    # Convert the image to RGB (opencv loads images in BGR format)
    rgb_image = cv2.cvtColor(image, cv2.COLOR_BGR2RGB)
    
    # Encode the faces
    face_encodings = face_recognition.face_encodings(rgb_image, face_locations)
    
    return face_locations, face_encodings, image

from PIL import Image, ImageDraw

def display_image_with_boxes(image, face_locations):
    # Convert the image to RGB format (OpenCV loads images in BGR format)
    image = cv2.cvtColor(image, cv2.COLOR_BGR2RGB)
    
    # Convert the image to a PIL Image
    pil_image = Image.fromarray(image)
    draw = ImageDraw.Draw(pil_image)

    # Draw rectangles around faces
    for (top, right, bottom, left) in face_locations:
        draw.rectangle([(left, top), (right, bottom)], outline="green", width=2)

    # Display the image
    pil_image.show()


def organize_photos(directory):
    face_encodings_dict = {}  # To store encodings and their corresponding image paths
    processed_images = set()
    image_paths = [os.path.join(directory, f) for f in os.listdir(directory) if f.endswith(('.png', '.jpg', '.jpeg'))]

    # Create Excel sheet if it doesn't exist
    if not os.path.exists(excel_path):
        create_excel_sheet()

    for image_path in image_paths:
        if image_path in processed_images:
            continue

        face_locations, face_encodings, image = detect_and_recognize_faces(image_path)
        
        if face_locations:
            display_image_with_boxes(image, face_locations)
        
        for encoding in face_encodings:
            match = False
            for name, encodings in face_encodings_dict.items():
                matches = face_recognition.compare_faces(encodings, encoding, tolerance=0.6)
                if True in matches:
                    face_encodings_dict[name].append(encoding)
                    os.makedirs(os.path.join(directory, name), exist_ok=True)
                    copyfile(image_path, os.path.join(directory, name, os.path.basename(image_path)))
                    update_excel_sheet(name, image_path)
                    match = True
                    break
            
            if not match:
                new_name = input(f"Enter a name for the new detected face in {image_path}: ").strip()
                if new_name:
                    face_encodings_dict[new_name] = [encoding]
                    os.makedirs(os.path.join(directory, new_name), exist_ok=True)
                    copyfile(image_path, os.path.join(directory, new_name, os.path.basename(image_path)))
                    update_excel_sheet(new_name, image_path)

        processed_images.add(image_path)

# Example usage
organize_photos('album')
